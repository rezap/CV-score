"""Builds realistic CV files for tests (no binary fixtures in the repo).

All personal details are fictional.
"""

from __future__ import annotations

from pathlib import Path

import docx
from docx.document import Document as DocxDocument
from docx.oxml import parse_xml
from docx.shared import Pt
from PIL import Image
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen.canvas import Canvas

PAGE_W, PAGE_H = A4
SIDEBAR_X, MAIN_X = 40, 220


def _lines(c: Canvas, x: float, y: float, lines: list[tuple[str, str, float]]) -> float:
    for text, font, size in lines:
        if text:
            c.setFont(font, size)
            c.drawString(x, y, text)
        y -= size * 1.45
    return y


def _running_header_footer(c: Canvas, page: int, pages: int) -> None:
    c.setFont("Helvetica", 8)
    c.drawString(40, PAGE_H - 25, "Jane Example - Curriculum Vitae")
    c.drawString(PAGE_W - 100, 20, f"Page {page} of {pages}")


def two_column_cv_pdf(path: Path) -> Path:
    """Two pages: a name banner across the top, a sidebar on the left, experience on the right."""
    c = Canvas(str(path), pagesize=A4)
    _running_header_footer(c, 1, 2)
    c.setFont("Helvetica-Bold", 22)
    c.drawString(40, PAGE_H - 70, "Jane Example")
    c.setFont("Helvetica", 11)
    c.drawString(40, PAGE_H - 88, "Senior Data Engineer")

    body = "Helvetica"
    top = PAGE_H - 130
    _lines(
        c,
        SIDEBAR_X,
        top,
        [
            ("CONTACT", "Helvetica-Bold", 10),
            ("jane@example.com", body, 9),
            ("+44 20 7946 0958", body, 9),
            ("", body, 9),
            ("SKILLS", "Helvetica-Bold", 10),
            ("• Python", body, 9),
            ("• SQL and dbt", body, 9),
            ("• Apache Spark", body, 9),
            ("", body, 9),
            ("LANGUAGES", "Helvetica-Bold", 10),
            ("English (native)", body, 9),
            ("German (B2)", body, 9),
        ],
    )
    _lines(
        c,
        MAIN_X,
        top,
        [
            ("EXPERIENCE", "Helvetica-Bold", 10),
            ("Senior Data Engineer, Acme Analytics Ltd, 2019 - 2024", "Helvetica-Bold", 9),
            ("• Led a team of 5 engineers building the company's streaming", body, 9),
            ("  platform, cutting data latency from hours to minutes.", body, 9),
            ("• Introduced data-quality monitoring and incident manage-", body, 9),
            ("  ment, reducing pipeline failures by 40%.", body, 9),
            ("", body, 9),
            ("Data Engineer, Northwind Traders, 2016 - 2019", "Helvetica-Bold", 9),
            ("• Built batch ETL pipelines in Python and Airflow.", body, 9),
        ],
    )
    c.showPage()
    _running_header_footer(c, 2, 2)
    _lines(
        c,
        MAIN_X,
        PAGE_H - 70,
        [
            ("EDUCATION", "Helvetica-Bold", 10),
            ("MSc Computer Science, University of Exampleton, 2016", body, 9),
            ("BSc Mathematics, University of Exampleton, 2014", body, 9),
        ],
    )
    _lines(
        c,
        SIDEBAR_X,
        PAGE_H - 70,
        [
            ("CERTIFICATIONS", "Helvetica-Bold", 10),
            ("AWS Data Analytics", body, 9),
            ("Databricks Associate", body, 9),
        ],
    )
    c.save()
    return path


def single_column_pdf(path: Path) -> Path:
    c = Canvas(str(path), pagesize=A4)
    y = PAGE_H - 60
    paragraph = [
        "I am writing to apply for the Senior Data Engineer role. Over the past eight",
        "years I have designed and operated data platforms for analytics teams, most",
        "recently leading a team of five engineers at Acme Analytics.",
    ]
    y = _lines(c, 50, y, [("Dear Hiring Manager,", "Helvetica", 10), ("", "Helvetica", 10)])
    y = _lines(c, 50, y, [(t, "Helvetica", 10) for t in paragraph])
    y = _lines(c, 50, y - 10, [("Kind regards,", "Helvetica", 10)])
    c.save()
    return path


def scanned_pdf(path: Path) -> Path:
    """A page with only an image and no text layer, like a scanned CV."""
    c = Canvas(str(path), pagesize=A4)
    c.drawImage(ImageReader(Image.new("RGB", (200, 280), "white")), 50, 300, 400, 500)
    c.save()
    return path


# --- Word ------------------------------------------------------------------------

_TEXTBOX_RUN = """
<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
     xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"
     xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
     xmlns:v="urn:schemas-microsoft-com:vml">
  <mc:AlternateContent>
    <mc:Choice Requires="wps">
      <w:drawing><wps:wsp><wps:txbx><w:txbxContent>
        <w:p><w:r><w:t>{text}</w:t></w:r></w:p>
      </w:txbxContent></wps:txbx></wps:wsp></w:drawing>
    </mc:Choice>
    <mc:Fallback>
      <w:pict><v:shape><v:textbox><w:txbxContent>
        <w:p><w:r><w:t>{text}</w:t></w:r></w:p>
      </w:txbxContent></v:textbox></v:shape></w:pict>
    </mc:Fallback>
  </mc:AlternateContent>
</w:r>
"""


def add_text_box(doc: DocxDocument, text: str) -> None:
    """Anchor a text box (as Word stores it: modern shape + legacy VML fallback)."""
    p = doc.add_paragraph()
    p._p.append(parse_xml(_TEXTBOX_RUN.format(text=text)))


def cv_docx(path: Path) -> Path:
    doc = docx.Document()
    header = doc.sections[0].header
    header.paragraphs[0].text = "Jane Example | jane@example.com | +44 20 7946 0958"

    name = doc.add_paragraph()
    run = name.add_run("Jane Example")
    run.bold = True
    run.font.size = Pt(20)
    add_text_box(doc, "Skills: Python, SQL, Spark")

    doc.add_heading("Experience", level=1)
    doc.add_paragraph("Senior Data Engineer, Acme Analytics Ltd, 2019 - 2024")
    doc.add_paragraph("Led a team of 5 engineers on the streaming platform.", style="List Bullet")
    doc.add_paragraph("Reduced pipeline failures by 40%.", style="List Bullet")
    # Manual bullets after soft line breaks, as people often type them.
    doc.add_paragraph("Data Engineer, Northwind Traders\n• Built ETL in Python\n• Ran Airflow")

    edu = doc.add_paragraph()
    edu.add_run("EDUCATION:").bold = True  # a "heading" typed in Normal style
    doc.add_paragraph("MSc Computer Science, University of Exampleton, 2016")

    doc.add_heading("Languages", level=1)
    table = doc.add_table(rows=3, cols=2)
    for row, (lang, level) in zip(
        table.rows, [("Language", "Level"), ("English", "Native"), ("German", "B2")], strict=True
    ):
        row.cells[0].text = lang
        row.cells[1].text = level
    doc.save(str(path))
    return path


def layout_table_docx(path: Path) -> Path:
    """A two-column CV built from a one-row table, as many templates do."""
    doc = docx.Document()
    table = doc.add_table(rows=1, cols=2)
    left, right = table.rows[0].cells
    left.paragraphs[0].text = "SKILLS"
    for t in ["Python", "SQL", "Spark"]:
        left.add_paragraph(t)
    right.paragraphs[0].text = "EXPERIENCE"
    for t in [
        "Senior Data Engineer, Acme Analytics Ltd",
        "Led a team of 5 engineers.",
        "Cut data latency from hours to minutes.",
    ]:
        right.add_paragraph(t)
    doc.save(str(path))
    return path


def date_column_cv_pdf(path: Path) -> Path:
    """Dates in a narrow left column, each next to its entry on the right."""
    c = Canvas(str(path), pagesize=A4)
    c.setFont("Helvetica-Bold", 20)
    c.drawString(40, PAGE_H - 60, "Sam Placeholder")
    body = "Helvetica"
    y = PAGE_H - 100
    y = _lines(c, 40, y, [("Experience", "Helvetica-Bold", 12)])
    rows = [
        ("2019 - 2024", "Product Manager, Contoso Ltd"),
        ("", "Launched a pricing feature used by 30% of customers."),
        ("2015 - 2019", "Business Analyst, Fabrikam Inc"),
        ("", "Built the weekly sales forecast used by the board."),
    ]
    for date, entry in rows:
        if date:
            c.setFont(body, 9)
            c.drawString(40, y, date)
        c.setFont(body, 9)
        c.drawString(140, y, entry)
        y -= 14
    _lines(c, 40, y - 10, [("Education", "Helvetica-Bold", 12)])
    c.setFont(body, 9)
    c.drawString(40, y - 40, "2011 - 2015")
    c.drawString(140, y - 40, "BA Economics, University of Exampleton")
    c.save()
    return path


def word_style_cv_docx(path: Path) -> Path:
    """A CV as people make it in Word: title, built-in heading and list styles, long prose."""
    doc = docx.Document()
    doc.add_heading("Alex Sample", 0)
    doc.add_paragraph("alex.sample@example.org | +31 6 1234 5678 | Rotterdam")
    doc.add_heading("Profile", 1)
    doc.add_paragraph(
        "Data scientist with seven years of experience building forecasting and pricing "
        "models for retail and logistics companies. I enjoy turning messy operational data "
        "into decisions that teams actually use, and I have led small teams through the full "
        "lifecycle from discovery to production monitoring. Fluent in Dutch and English."
    )
    doc.add_heading("Experience", 1)
    doc.add_paragraph().add_run(
        "Lead Data Scientist, Example Logistics BV, 2020 - present"
    ).bold = True
    for text in [
        "Designed a demand forecasting system that reduced stock-outs by 18% across 120 "
        "warehouses while keeping inventory costs flat.",
        "Hired and mentored four data scientists; introduced code review, experiment tracking "
        "and model monitoring practices.",
    ]:
        doc.add_paragraph(text, style="List Bullet")
    doc.add_heading("Education", 1)
    doc.add_paragraph("MSc Econometrics, Erasmus University Rotterdam, 2017")
    doc.add_heading("Skills", 1)
    for text in ["Python, SQL, Spark", "Forecasting, causal inference"]:
        doc.add_paragraph(text, style="List Bullet")
    doc.save(str(path))
    return path


def marketing_cv_docx(path: Path) -> Path:
    """A clearly non-matching CV for a data engineering role."""
    doc = docx.Document()
    doc.add_heading("Morgan Lee", 0)
    doc.add_paragraph("morgan.lee@example.net | +44 161 496 0000 | Manchester")
    doc.add_heading("Profile", 1)
    doc.add_paragraph(
        "Marketing coordinator with four years of experience planning events and social "
        "media campaigns for consumer brands. I love telling stories and working with people."
    )
    doc.add_heading("Experience", 1)
    doc.add_paragraph().add_run(
        "Marketing Coordinator, Bright Events Ltd, 2021 - present"
    ).bold = True
    for text in [
        "Organised 25 trade-show events per year with budgets up to 50,000 pounds.",
        "Grew the company's Instagram following from 2,000 to 15,000 in two years.",
    ]:
        doc.add_paragraph(text, style="List Bullet")
    doc.add_heading("Education", 1)
    doc.add_paragraph("BA Communications, University of Exampleton, 2020")
    doc.add_heading("Skills", 1)
    for text in ["Canva, Hootsuite, Mailchimp", "Event planning, copywriting"]:
        doc.add_paragraph(text, style="List Bullet")
    doc.save(str(path))
    return path


def _words_without_spaces(
    c: Canvas, x: float, y: float, text: str, font: str, size: float, gap: float
) -> None:
    """Draw each word on its own with a gap of ``gap`` em and no space characters, as LaTeX
    does; tight justification makes the gaps smaller than a fixed 3 pt tolerance."""
    c.setFont(font, size)
    for word in text.split():
        c.drawString(x, y, word)
        x += c.stringWidth(word, font, size) + gap * size


def _letter_spaced(c: Canvas, x: float, y: float, text: str, font: str, size: float) -> None:
    """Tracked heading: every letter placed on its own, 0.35 em apart, wider at word spaces."""
    c.setFont(font, size)
    for ch in text:
        if ch != " ":
            c.drawString(x, y, ch)
        x += c.stringWidth(ch, font, size) + 0.35 * size


def latex_style_pdf(path: Path) -> Path:
    """Small, tightly justified text with no space characters (word gaps of 0.15 em)."""
    c = Canvas(str(path), pagesize=A4)
    y = PAGE_H - 60
    for text, font, size in [
        ("Erik Lindqvist", "Helvetica-Bold", 16),
        ("Experience", "Helvetica-Bold", 12),
        (
            "Built an internal app to run batch campaigns for marketing and analytics teams.",
            "Helvetica",
            8,
        ),
        (
            "Reduced delivery time by thirty percent through queuing and A/B testing.",
            "Helvetica",
            8,
        ),
    ]:
        _words_without_spaces(c, 60, y, text, font, size, gap=0.15)
        y -= size * 1.8
    c.showPage()
    c.save()
    return path


def design_tool_pdf(path: Path) -> Path:
    """A design-tool layout: narrow sidebar, main column ending at two thirds of the page
    (a wide empty right margin), a big name wrapped over two lines, tracked headings."""
    c = Canvas(str(path), pagesize=A4)
    c.setFont("Helvetica-Bold", 30)
    c.drawString(30, PAGE_H - 60, "SOFIE")
    c.drawString(30, PAGE_H - 95, "HANSEN")
    side, main = 30, 210
    y = PAGE_H - 160
    _letter_spaced(c, side, y, "CONTACT", "Helvetica-Bold", 9)
    _letter_spaced(c, main, y, "ABOUT ME", "Helvetica-Bold", 9)
    y = _lines(
        c,
        main,
        y - 18,
        [
            ("Curious, data-driven marketer who loves", "Helvetica", 9),
            ("building brands people talk about.", "Helvetica", 9),
        ],
    )
    _lines(
        c,
        side,
        PAGE_H - 178,
        [
            ("sofie.hansen@example.com", "Helvetica", 8),
            ("+47 98 76 54 32", "Helvetica", 8),
            ("Oslo, Norway", "Helvetica", 8),
        ],
    )
    y -= 14
    _letter_spaced(c, side, y, "SKILLS", "Helvetica-Bold", 9)
    _letter_spaced(c, main, y, "EXPERIENCE", "Helvetica-Bold", 9)
    _lines(
        c,
        side,
        y - 18,
        [
            ("SEO and SEM", "Helvetica", 8),
            ("Paid social", "Helvetica", 8),
            ("Copywriting", "Helvetica", 8),
        ],
    )
    _lines(
        c,
        main,
        y - 18,
        [
            ("Marketing Manager, Nordic Coffee, 2019 - 2024", "Helvetica-Bold", 9),
            ("Grew online sales from NOK 4M to NOK 11M in", "Helvetica", 9),
            ("three years and launched a subscription.", "Helvetica", 9),
            ("", "Helvetica", 9),
            ("Marketing Specialist, Oslo Media, 2016 - 2019", "Helvetica-Bold", 9),
            ("Ran paid social campaigns for 25 clients.", "Helvetica", 9),
        ],
    )
    c.showPage()
    c.save()
    return path


def password_pdf(path: Path) -> Path:
    from reportlab.lib import pdfencrypt

    c = Canvas(str(path), pagesize=A4, encrypt=pdfencrypt.StandardEncryption("secret"))
    c.drawString(60, 700, "Jane Example")
    c.showPage()
    c.save()
    return path


def scanned_cv_image(path: Path) -> Path:
    """A photo/scan of a CV page as an image file (PNG or JPEG by suffix)."""
    from PIL import ImageDraw

    img = Image.new("RGB", (850, 1100), "white")
    draw = ImageDraw.Draw(img)
    for i, line in enumerate(["Kari Nordmann", "Experience", "Data Scientist, Equinor"]):
        draw.text((60, 60 + i * 40), line, fill="black")
    img.save(path)
    return path


def text_then_scan_pdf(path: Path) -> Path:
    """Page 1 has a text layer; page 2 is a scanned image (e.g. a scanned diploma page)."""
    c = Canvas(str(path), pagesize=A4)
    _lines(
        c,
        60,
        PAGE_H - 80,
        [
            ("Kari Nordmann", "Helvetica-Bold", 18),
            ("Experience", "Helvetica-Bold", 13),
            (
                "Senior Data Scientist, Equinor, 2019 - present. Built forecasting models for",
                "Helvetica",
                10,
            ),
            ("offshore platforms and led a team of four engineers.", "Helvetica", 10),
            (
                "Data Analyst, Statkraft, 2015 - 2019. Automated reporting with Python and SQL.",
                "Helvetica",
                10,
            ),
            ("Education", "Helvetica-Bold", 13),
            ("MSc Statistics, NTNU, 2013 - 2015", "Helvetica", 10),
        ],
    )
    c.showPage()
    c.drawImage(ImageReader(Image.new("RGB", (200, 280), "white")), 50, 300, 400, 500)
    c.showPage()
    c.save()
    return path


def readable_scan_pdf(path: Path, lines: list[str]) -> Path:
    """A scanned CV: the page is one image of printed text, with no text layer."""
    from PIL import ImageDraw, ImageFont

    img = Image.new("RGB", (1240, 1754), "white")  # A4 at 150 dpi
    draw = ImageDraw.Draw(img)
    y = 120
    for line in lines:
        heading = line.isupper()
        font = ImageFont.load_default(size=40 if heading else 28)
        if line:
            draw.text((120, y), line, fill="black", font=font)
        y += 60 if heading else 44
    c = Canvas(str(path), pagesize=A4)
    c.drawImage(ImageReader(img), 0, 0, PAGE_W, PAGE_H)
    c.showPage()
    c.save()
    return path


def scan_with_text_footer_pdf(path: Path) -> Path:
    """Page 1 has text; page 2 is a full-page scan with a real-text page number on top."""
    text_then_scan_pdf(path)
    c = Canvas(str(path), pagesize=A4)
    _lines(
        c,
        60,
        PAGE_H - 80,
        [
            ("Kari Nordmann", "Helvetica-Bold", 18),
            ("Experience", "Helvetica-Bold", 13),
            (
                "Senior Data Scientist, Equinor, 2019 - present. Built forecasting models for",
                "Helvetica",
                10,
            ),
            ("offshore platforms and led a team of four engineers.", "Helvetica", 10),
            (
                "Data Analyst, Statkraft, 2015 - 2019. Automated reporting with Python and SQL.",
                "Helvetica",
                10,
            ),
        ],
    )
    c.showPage()
    c.drawImage(ImageReader(Image.new("RGB", (200, 280), "white")), 0, 0, PAGE_W, PAGE_H)
    c.setFont("Helvetica", 8)
    c.drawString(PAGE_W / 2, 20, "Page 2 of 2")
    c.showPage()
    c.save()
    return path
