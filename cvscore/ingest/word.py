"""Word (.docx) text extraction with python-docx.

Handles what CV templates commonly do:
- paragraphs and tables in document order, with heading and list styles;
- tables used for page layout (two-column CVs) are read cell by cell, data tables row by row;
- text boxes (sidebars), which python-docx does not expose, are read from the XML;
- headers and footers, which often hold contact details, are flagged as contact text.
"""

from __future__ import annotations

from collections.abc import Iterator
from pathlib import Path

import docx
from docx.document import Document as DocxDocument
from docx.opc.exceptions import PackageNotFoundError
from docx.oxml.ns import qn
from docx.oxml.xmlchemy import BaseOxmlElement
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

from cvscore.ingest.document import ExtractionError, build_document
from cvscore.ingest.text_utils import is_section_keyword, join_lines, normalize_text, strip_bullet
from cvscore.models import BlockKind, DocType, SourceDocument, TextBlock

_MC_FALLBACK = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"
_HEADING_STYLES = ("heading", "title", "überschrift", "titel")
_LIST_STYLES = ("list", "aufzählung", "liste")
LAYOUT_CELL_PARAGRAPHS = 3  # a cell with this many paragraphs means a layout table


def _style_name(p: Paragraph) -> str:
    try:
        return (p.style.name or "").lower() if p.style is not None else ""
    except (KeyError, ValueError):
        return ""


def _is_list(p: Paragraph) -> bool:
    ppr = p._p.pPr
    has_numbering = ppr is not None and ppr.numPr is not None
    return has_numbering or any(s in _style_name(p) for s in _LIST_STYLES)


def _text_boxes(el: BaseOxmlElement) -> Iterator[BaseOxmlElement]:
    """``w:txbxContent`` elements under ``el``, skipping legacy VML duplicates.

    Word stores modern text boxes twice (``mc:Choice`` + ``mc:Fallback``); reading only
    the non-fallback copy avoids duplicate text.
    """
    for box in el.iter(qn("w:txbxContent")):
        if not any(a.tag == _MC_FALLBACK for a in box.iterancestors()):
            yield box


def _paragraph_blocks(p: Paragraph) -> list[TextBlock]:
    raw = p.text
    blocks: list[TextBlock] = []
    if raw.strip():
        style = _style_name(p)
        lines = [normalize_text(x) for x in raw.splitlines()]
        lines = [x for x in lines if x]
        single = len(lines) == 1
        if single and (style.startswith(_HEADING_STYLES) or is_section_keyword(lines[0])):
            blocks.append(TextBlock(kind=BlockKind.HEADING, text=lines[0].rstrip(":")))
        elif _is_list(p):
            _, text = strip_bullet(join_lines(lines))
            blocks.append(TextBlock(kind=BlockKind.LIST_ITEM, text=text))
        else:
            # Soft line breaks inside one paragraph are common; manual bullets after them
            # become separate list items.
            buf: list[str] = []
            for line in lines:
                bullet, text = strip_bullet(line)
                if bullet:
                    if buf:
                        blocks.append(TextBlock(kind=BlockKind.PARAGRAPH, text=join_lines(buf)))
                        buf = []
                    blocks.append(TextBlock(kind=BlockKind.LIST_ITEM, text=text))
                elif single and is_section_keyword(text):
                    blocks.append(TextBlock(kind=BlockKind.HEADING, text=text.rstrip(":")))
                else:
                    buf.append(text)
            if buf:
                blocks.append(TextBlock(kind=BlockKind.PARAGRAPH, text=join_lines(buf)))
    for box in _text_boxes(p._p):
        for child in box.iterchildren(qn("w:p")):
            blocks.extend(_paragraph_blocks(Paragraph(child, p._parent)))
    return blocks


def _unique_cells(table: Table) -> Iterator[list[_Cell]]:
    """Rows of cells with horizontally merged cells reported once."""
    for row in table.rows:
        seen: set[int] = set()
        cells = []
        for cell in row.cells:
            if id(cell._tc) not in seen:
                seen.add(id(cell._tc))
                cells.append(cell)
        yield cells


def _is_layout_table(table: Table) -> bool:
    for cells in _unique_cells(table):
        for cell in cells:
            if sum(1 for p in cell.paragraphs if p.text.strip()) >= LAYOUT_CELL_PARAGRAPHS:
                return True
            if cell.tables:
                return True
    return False


def _table_blocks(table: Table) -> list[TextBlock]:
    blocks: list[TextBlock] = []
    if _is_layout_table(table):
        # Read like a page: column by column within each row.
        for cells in _unique_cells(table):
            for cell in cells:
                blocks.extend(_container_blocks(cell))
        return blocks
    for cells in _unique_cells(table):
        texts = [normalize_text(c.text.replace("\n", " ")) for c in cells]
        texts = [t for t in texts if t]
        if texts:
            blocks.append(TextBlock(kind=BlockKind.TABLE_ROW, text=" | ".join(texts)))
    return blocks


def _container_blocks(container: DocxDocument | _Cell) -> list[TextBlock]:
    blocks: list[TextBlock] = []
    for item in container.iter_inner_content():
        if isinstance(item, Paragraph):
            blocks.extend(_paragraph_blocks(item))
        else:
            blocks.extend(_table_blocks(item))
    return blocks


def _header_footer_blocks(doc: DocxDocument) -> list[TextBlock]:
    blocks: list[TextBlock] = []
    seen: set[str] = set()
    for section in doc.sections:
        for part in (section.header, section.first_page_header, section.footer):
            if part.is_linked_to_previous:
                continue
            for item in part.iter_inner_content():
                found = (
                    _paragraph_blocks(item) if isinstance(item, Paragraph) else _table_blocks(item)
                )
                for b in found:
                    if b.text not in seen:
                        seen.add(b.text)
                        b.is_contact = True
                        b.section = "Header/footer"
                        blocks.append(b)
    return blocks


def extract_docx(path: Path, *, doc_type: DocType, display_path: str) -> SourceDocument:
    try:
        doc = docx.Document(str(path))
    except PackageNotFoundError as e:
        # python-docx says "Package not found" for files that exist but are not valid .docx.
        raise ExtractionError(
            "Not a valid .docx file (corrupt, or another format renamed to .docx)."
        ) from e
    except Exception as e:
        raise ExtractionError(f"Cannot open Word document: {e}") from e
    body = _container_blocks(doc)
    return build_document(
        doc_type=doc_type,
        path=path,
        display_path=display_path,
        blocks=[*_header_footer_blocks(doc), *body],
        extractor="python-docx",
    )
